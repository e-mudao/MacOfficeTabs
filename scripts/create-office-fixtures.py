"""Disposable local samples for window/focus tests; never changes existing files."""
from pathlib import Path
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

root = Path(__file__).resolve().parent.parent / 'docs/evidence'
for app, extension in [('excel','xlsx'),('powerpoint','pptx')]:
    for index, relative in enumerate([f'a/same.{extension}',f'b/same.{extension}',f'test-03.{extension}'],1):
        target = root / f'fixtures-{app}' / relative
        target.parent.mkdir(parents=True,exist_ok=True)
        if target.exists():
            print('Preserved:',target.relative_to(root)); continue
        if app == 'excel':
            book = Workbook(); sheet = book.active
            sheet['A1'] = f'MacOfficeTabs disposable Excel fixture {index}'
            sheet['A2'] = 'Keyboard test: type in B2 and undo.'
            sheet.column_dimensions['A'].width = 52
            book.save(target); load_workbook(target).close()
        else:
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(0.6),Inches(0.8),Inches(8),Inches(3))
            box.text_frame.text = f'MacOfficeTabs disposable PowerPoint fixture {index}\nKeyboard test: type here and undo.'
            deck.save(target); Presentation(target)
        print('Created and reopened:',target.relative_to(root))
